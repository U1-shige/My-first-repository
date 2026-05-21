#!/usr/bin/env python3
"""Generate meeting minutes as .docx from a JSON data file."""
import sys
import json
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def add_label_value(doc, label, value):
    p = doc.add_paragraph()
    run_label = p.add_run(f"{label}：")
    run_label.bold = True
    p.add_run(value or "（要確認）")


def add_section_heading(doc, text, level=2):
    doc.add_heading(text, level=level)


def create_minutes(data: dict, output_path: str):
    doc = Document()

    # Title
    title = doc.add_heading("議事録", 0)

    # Basic info
    doc.add_paragraph()
    add_label_value(doc, "会議名", data.get("meeting_name"))
    add_label_value(doc, "日時", data.get("datetime"))
    add_label_value(doc, "場所", data.get("location"))
    add_label_value(doc, "参加者", data.get("participants"))
    doc.add_paragraph()

    # Agenda
    add_section_heading(doc, "議題")
    for i, item in enumerate(data.get("agenda", []), 1):
        doc.add_paragraph(f"{i}　{item}")
    doc.add_paragraph()

    # Discussion
    add_section_heading(doc, "議事内容")
    for i, disc in enumerate(data.get("discussions", []), 1):
        add_section_heading(doc, f"{i}　{disc.get('topic', '')}", level=3)

        p = doc.add_paragraph()
        p.add_run("概要：").bold = True
        doc.add_paragraph(disc.get("summary", ""))

        p = doc.add_paragraph()
        p.add_run("決定事項：").bold = True
        doc.add_paragraph(disc.get("decisions", ""))

        p = doc.add_paragraph()
        p.add_run("保留事項：").bold = True
        doc.add_paragraph(disc.get("pending", ""))
        doc.add_paragraph()

    # Flow summary
    add_section_heading(doc, "協議の流れ（概要）")
    doc.add_paragraph(data.get("flow_summary", ""))

    doc.save(output_path)
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python3 create_minutes_docx.py <data.json> <output.docx>")
        sys.exit(1)
    with open(sys.argv[1], encoding="utf-8") as f:
        data = json.load(f)
    create_minutes(data, sys.argv[2])
